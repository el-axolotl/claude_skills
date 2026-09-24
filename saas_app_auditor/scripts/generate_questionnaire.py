#!/usr/bin/env python3
"""Build the SaaS security audit questionnaire as a formatted .xlsx file.

Input is a JSON file (see saas_app_auditor/assets/questionnaire_schema.json for the
shape) containing the app name and a list of questionnaire rows already researched
and written by Claude. This script only handles layout/formatting of the workbook —
it does no research and invents no content.

Usage:
    python generate_questionnaire.py --input rows.json --output "Acme_Questionnaire.xlsx"
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit(
        "openpyxl is required but not installed.\n"
        "Run: pip install openpyxl"
    )

COLUMNS = [
    "MITRE Category",
    "Severity",
    "Question",
    "Answer",
    "Description/Abuse Context",
]

SEVERITY_LEVELS = ["Critical", "High", "Medium", "Low"]

SEVERITY_FILL = {
    "Critical": "C00000",
    "High": "E97132",
    "Medium": "FFC000",
    "Low": "70AD47",
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

COLUMN_WIDTHS = {
    "MITRE Category": 34,
    "Severity": 12,
    "Question": 50,
    "Answer": 30,
    "Description/Abuse Context": 55,
}


def sanitize_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "", name).strip()
    return (name or "Questionnaire")[:31]


def load_rows(data: dict) -> list:
    rows = data.get("rows", [])
    if not rows:
        sys.exit("Input JSON has no rows to write.")

    required = {"mitre_category", "severity", "question", "description"}
    for i, row in enumerate(rows, start=1):
        missing = required - row.keys()
        if missing:
            sys.exit(f"Row {i} is missing field(s): {', '.join(sorted(missing))}")
        severity = row["severity"].strip().title()
        if severity not in SEVERITY_LEVELS:
            sys.exit(
                f"Row {i} has invalid severity '{row['severity']}'. "
                f"Must be one of: {', '.join(SEVERITY_LEVELS)}"
            )
        row["severity"] = severity
    return rows


def build_workbook(app_name: str, rows: list) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_name(f"{app_name} Audit")

    ws.append(COLUMNS)
    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = THIN_BORDER

    for row in rows:
        ws.append(
            [
                row["mitre_category"],
                row["severity"],
                row["question"],
                row.get("answer", ""),
                row["description"],
            ]
        )

    last_row = ws.max_row
    for r in range(2, last_row + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = WRAP_TOP_LEFT
            cell.border = THIN_BORDER
        severity_cell = ws.cell(row=r, column=2)
        fill_color = SEVERITY_FILL.get(severity_cell.value)
        if fill_color:
            severity_cell.fill = PatternFill("solid", fgColor=fill_color)
            severity_cell.font = Font(color="FFFFFF", bold=True)
            severity_cell.alignment = Alignment(vertical="top", horizontal="center")

    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS[header]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(SEVERITY_LEVELS)}"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    dv.add(f"B2:B{last_row}")

    ws.row_dimensions[1].height = 30

    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to the rows JSON file")
    parser.add_argument(
        "--output",
        help="Path to write the .xlsx file (defaults to '<AppName>_Security_Questionnaire.xlsx')",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        sys.exit(f"Input file not found: {input_path}")

    data = json.loads(input_path.read_text(encoding="utf-8"))
    app_name = data.get("app_name", "").strip()
    if not app_name:
        sys.exit("Input JSON must include a non-empty 'app_name'.")

    rows = load_rows(data)
    wb = build_workbook(app_name, rows)

    output_path = Path(
        args.output or f"{re.sub(r'[^A-Za-z0-9_-]+', '_', app_name)}_Security_Questionnaire.xlsx"
    )
    wb.save(output_path)
    print(f"Wrote {len(rows)} questions to {output_path}")


if __name__ == "__main__":
    main()
