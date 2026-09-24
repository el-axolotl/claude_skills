#!/usr/bin/env python3
"""Build the STRIDE threat model as a formatted .xlsx file.

Input is a JSON file (see assets/threat_model_schema_example.json for the shape)
containing the app name and a list of threat rows already derived from a filled-out
questionnaire. This script only handles layout/formatting of the workbook — it does
no threat analysis and invents no content.

Usage:
    python generate_threat_model.py --input threats.json --output "Acme_Threat_Model.xlsx"
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit(
        "openpyxl is required but not installed.\n"
        "Run: pip install openpyxl"
    )

COLUMNS = [
    "STRIDE Category",
    "System Component",
    "Threat",
    "Current Mitigation",
    "Mitigation Status",
    "Linked Tickets",
]

STRIDE_CATEGORIES = [
    "Spoofing",
    "Tampering",
    "Repudiation",
    "Information Disclosure",
    "Denial of Service",
    "Elevation of Privilege",
]

MITIGATION_STATUSES = [
    "Not Mitigated",
    "Partially Mitigated",
    "Mitigated",
    "Risk Accepted",
    "Unknown",
]

STATUS_FILL = {
    "Not Mitigated": "C00000",
    "Partially Mitigated": "E97132",
    "Mitigated": "70AD47",
    "Risk Accepted": "4472C4",
    "Unknown": "A6A6A6",
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

COLUMN_WIDTHS = {
    "STRIDE Category": 20,
    "System Component": 26,
    "Threat": 45,
    "Current Mitigation": 40,
    "Mitigation Status": 18,
    "Linked Tickets": 18,
}


def sanitize_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "", name).strip()
    return (name or "Threat Model")[:31]


def _canonicalize(value: str, options: list, field_name: str, row_num: int) -> str:
    lookup = {opt.lower(): opt for opt in options}
    canonical = lookup.get(value.strip().lower())
    if canonical is None:
        sys.exit(
            f"Row {row_num} has invalid {field_name} '{value}'. "
            f"Must be one of: {', '.join(options)}"
        )
    return canonical


def load_rows(data: dict) -> list:
    rows = data.get("rows", [])
    if not rows:
        sys.exit("Input JSON has no rows to write.")

    required = {"stride_category", "system_component", "threat", "current_mitigation", "mitigation_status"}
    for i, row in enumerate(rows, start=1):
        missing = required - row.keys()
        if missing:
            sys.exit(f"Row {i} is missing field(s): {', '.join(sorted(missing))}")

        row["stride_category"] = _canonicalize(row["stride_category"], STRIDE_CATEGORIES, "stride_category", i)
        row["mitigation_status"] = _canonicalize(row["mitigation_status"], MITIGATION_STATUSES, "mitigation_status", i)
    return rows


def build_workbook(app_name: str, rows: list) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_name(f"{app_name} Threat Model")

    ws.append(COLUMNS)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = THIN_BORDER

    for row in rows:
        ws.append(
            [
                row["stride_category"],
                row["system_component"],
                row["threat"],
                row["current_mitigation"],
                row["mitigation_status"],
                row.get("linked_tickets", ""),
            ]
        )

    last_row = ws.max_row
    for r in range(2, last_row + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = WRAP_TOP_LEFT
            cell.border = THIN_BORDER
        status_cell = ws.cell(row=r, column=5)
        fill_color = STATUS_FILL.get(status_cell.value)
        if fill_color:
            status_cell.fill = PatternFill("solid", fgColor=fill_color)
            status_cell.font = Font(color="FFFFFF", bold=True)
            status_cell.alignment = Alignment(vertical="top", horizontal="center")

    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS[header]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    stride_dv = DataValidation(
        type="list",
        formula1=f'"{",".join(STRIDE_CATEGORIES)}"',
        allow_blank=True,
    )
    ws.add_data_validation(stride_dv)
    stride_dv.add(f"A2:A{last_row}")

    status_dv = DataValidation(
        type="list",
        formula1=f'"{",".join(MITIGATION_STATUSES)}"',
        allow_blank=True,
    )
    ws.add_data_validation(status_dv)
    status_dv.add(f"E2:E{last_row}")

    ws.row_dimensions[1].height = 30

    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to the threat rows JSON file")
    parser.add_argument(
        "--output",
        help="Path to write the .xlsx file (defaults to '<AppName>_Threat_Model.xlsx')",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        sys.exit(f"Input file not found: {input_path}")

    data = json.loads(input_path.read_text(encoding="utf-8"))
    app_name = data.get("app_name", "").strip()
    if not app_name:
        sys.exit("Input JSON must include a non-empty 'app_name'.")

    rows = load_rows(data)
    wb = build_workbook(app_name, rows)

    output_path = Path(
        args.output or f"{re.sub(r'[^A-Za-z0-9_-]+', '_', app_name)}_Threat_Model.xlsx"
    )
    wb.save(output_path)
    print(f"Wrote {len(rows)} threats to {output_path}")


if __name__ == "__main__":
    main()
